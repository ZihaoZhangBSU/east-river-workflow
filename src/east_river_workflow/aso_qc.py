"""ASO quality control, diagnostics, cleaned rasters, and a Word report.

The quality-control policy is deliberately conservative. The original ASO
GeoTIFFs are never modified. Only values that are physically impossible under
explicit, user-configurable hard limits are changed in the cleaned copies.
High-but-allowed values are reported for review and retained.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
import json
import logging
from pathlib import Path
from typing import Any

import geopandas as gpd
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import rasterio
from matplotlib.colors import ListedColormap
from rasterio.features import rasterize
from shapely.geometry import mapping

from .config import WorkflowConfig
from .data_access import read_aso_raster
from .grids import GridSpec
from .utils import save_json


VARIABLE_SLUGS = {"SWE": "swe", "Snow depth": "snow_depth"}


@dataclass
class ASOQCResult:
    """One raster after ASO quality control."""

    date: str
    variable: str
    source_path: Path
    grid: GridSpec
    metadata: dict[str, str]
    raw_mm: np.ndarray
    cleaned_mm: np.ndarray
    inside_watershed: np.ndarray
    clipped_to_zero: np.ndarray
    masked_negative: np.ndarray
    masked_above_maximum: np.ndarray
    hard_maximum_mm: float
    summary: dict[str, Any]

    @property
    def adjusted_mask(self) -> np.ndarray:
        return self.clipped_to_zero | self.masked_negative | self.masked_above_maximum


def _percentile_stats(values: np.ndarray, prefix: str) -> dict[str, float | int]:
    data = np.asarray(values, dtype=float)
    data = data[np.isfinite(data)]
    if data.size == 0:
        return {
            f"{prefix}_count": 0,
            f"{prefix}_min_mm": np.nan,
            f"{prefix}_p50_mm": np.nan,
            f"{prefix}_p95_mm": np.nan,
            f"{prefix}_p99_mm": np.nan,
            f"{prefix}_p99_9_mm": np.nan,
            f"{prefix}_p99_99_mm": np.nan,
            f"{prefix}_max_mm": np.nan,
        }
    percentiles = np.percentile(data, [50, 95, 99, 99.9, 99.99])
    return {
        f"{prefix}_count": int(data.size),
        f"{prefix}_min_mm": float(np.min(data)),
        f"{prefix}_p50_mm": float(percentiles[0]),
        f"{prefix}_p95_mm": float(percentiles[1]),
        f"{prefix}_p99_mm": float(percentiles[2]),
        f"{prefix}_p99_9_mm": float(percentiles[3]),
        f"{prefix}_p99_99_mm": float(percentiles[4]),
        f"{prefix}_max_mm": float(np.max(data)),
    }


def _hard_maximum(variable: str, qc: dict[str, Any]) -> float:
    if variable == "SWE":
        return float(qc["swe_hard_max_mm"])
    if variable == "Snow depth":
        return float(qc["snow_depth_hard_max_mm"])
    raise ValueError(f"Unsupported ASO variable: {variable}")


def watershed_center_mask(watershed: gpd.GeoDataFrame, grid: GridSpec) -> np.ndarray:
    """Return a center-based watershed mask for QC counts and maps.

    The scientific regridding and basin calculations still use the exact
    fractional-overlap weights. This center mask is only a diagnostic aid.
    """
    local = watershed.to_crs(grid.crs)
    geometry = local.geometry.union_all()
    return rasterize(
        [(mapping(geometry), 1)],
        out_shape=grid.shape,
        transform=grid.transform,
        fill=0,
        all_touched=False,
        dtype="uint8",
    ).astype(bool)


def apply_aso_quality_control(
    raw_mm: np.ndarray,
    *,
    date: str,
    variable: str,
    source_path: str | Path,
    grid: GridSpec,
    metadata: dict[str, str],
    inside_watershed: np.ndarray,
    qc: dict[str, Any],
) -> ASOQCResult:
    """Apply conservative, explicit ASO quality-control rules.

    Rules:
    1. Existing no-data and non-finite values remain no-data.
    2. Tiny negative values within ``negative_clip_tolerance_mm`` are set to 0.
    3. More-negative values are set to no-data.
    4. Values greater than the variable-specific hard maximum are set to no-data.

    No percentile clipping and no model-based correction are performed.
    """
    raw = np.asarray(raw_mm, dtype=float)
    if raw.shape != grid.shape:
        raise ValueError(f"ASO array shape {raw.shape} does not match grid shape {grid.shape}.")
    if inside_watershed.shape != raw.shape:
        raise ValueError("Watershed diagnostic mask does not match the ASO grid.")

    hard_maximum = _hard_maximum(variable, qc)
    tolerance = abs(float(qc["negative_clip_tolerance_mm"]))
    finite = np.isfinite(raw)
    clipped_to_zero = finite & (raw < 0.0) & (raw >= -tolerance)
    masked_negative = finite & (raw < -tolerance)
    masked_above = finite & (raw > hard_maximum)

    cleaned = raw.copy()
    cleaned[clipped_to_zero] = 0.0
    cleaned[masked_negative | masked_above] = np.nan
    adjusted = clipped_to_zero | masked_negative | masked_above

    raw_count = int(np.count_nonzero(finite))
    adjusted_count = int(np.count_nonzero(adjusted))
    inside_valid = finite & inside_watershed
    inside_adjusted = adjusted & inside_watershed
    summary: dict[str, Any] = {
        "date": date,
        "variable": variable,
        "source_path": str(Path(source_path).expanduser()),
        "hard_maximum_mm": hard_maximum,
        "negative_clip_tolerance_mm": tolerance,
        "nodata_or_nonfinite_count": int(raw.size - raw_count),
        "clipped_to_zero_count": int(np.count_nonzero(clipped_to_zero)),
        "masked_negative_count": int(np.count_nonzero(masked_negative)),
        "masked_above_hard_max_count": int(np.count_nonzero(masked_above)),
        "adjusted_count": adjusted_count,
        "adjusted_percent_of_raw_valid": (100.0 * adjusted_count / raw_count) if raw_count else np.nan,
        "watershed_raw_valid_count": int(np.count_nonzero(inside_valid)),
        "watershed_adjusted_count": int(np.count_nonzero(inside_adjusted)),
        "watershed_adjusted_percent": (
            100.0 * np.count_nonzero(inside_adjusted) / np.count_nonzero(inside_valid)
            if np.count_nonzero(inside_valid)
            else np.nan
        ),
    }
    summary.update(_percentile_stats(raw, "raw"))
    summary.update(_percentile_stats(cleaned, "cleaned"))
    summary.update(_percentile_stats(raw[inside_watershed], "watershed_raw"))
    summary.update(_percentile_stats(cleaned[inside_watershed], "watershed_cleaned"))

    return ASOQCResult(
        date=date,
        variable=variable,
        source_path=Path(source_path).expanduser(),
        grid=grid,
        metadata=metadata,
        raw_mm=raw,
        cleaned_mm=cleaned,
        inside_watershed=inside_watershed,
        clipped_to_zero=clipped_to_zero,
        masked_negative=masked_negative,
        masked_above_maximum=masked_above,
        hard_maximum_mm=hard_maximum,
        summary=summary,
    )


def qc_cleaned_path(cfg: WorkflowConfig, date: str, variable: str) -> Path:
    slug = VARIABLE_SLUGS[variable]
    return cfg.output_dir / "cache" / "aso_qc" / f"{date}_{slug}_qc_mm.tif"


def _write_cleaned_raster(path: Path, result: ASOQCResult, qc: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    grid = result.grid
    nodata = -9999.0
    with rasterio.open(
        path,
        "w",
        driver="GTiff",
        width=grid.width,
        height=grid.height,
        count=1,
        dtype="float32",
        crs=grid.crs,
        transform=grid.transform,
        nodata=nodata,
        compress="deflate",
        predictor=3,
    ) as dst:
        dst.write(np.where(np.isfinite(result.cleaned_mm), result.cleaned_mm, nodata).astype(np.float32), 1)
        dst.set_band_description(1, f"ASO {result.variable} after QC [mm]")
        dst.update_tags(
            units="mm",
            source_file=str(result.source_path),
            qc_action="mask hard-invalid values; clip tiny negatives to zero",
            qc_hard_maximum_mm=str(result.hard_maximum_mm),
            qc_negative_clip_tolerance_mm=str(qc["negative_clip_tolerance_mm"]),
            qc_adjusted_count=str(result.summary["adjusted_count"]),
        )


def read_cleaned_aso_raster(path: str | Path) -> tuple[np.ndarray, GridSpec, dict[str, str]]:
    """Read a cleaned ASO raster whose stored units are millimeters."""
    from .grids import grid_from_centers

    with rasterio.open(path) as src:
        data = src.read(1, masked=True).astype(float).filled(np.nan)
        tags = src.tags()
        units = str(tags.get("units", "")).strip().lower()
        if units not in {"mm", "millimeter", "millimeters", "millimetre", "millimetres"}:
            raise ValueError(f"Cleaned ASO raster {path} must declare millimeter units; tags={tags}")
        x = src.transform.c + (np.arange(src.width) + 0.5) * src.transform.a
        y = src.transform.f + (np.arange(src.height) + 0.5) * src.transform.e
        grid = grid_from_centers(f"ASO-QC:{Path(path).stem}", x, y, src.crs)
        metadata = {
            "tags": str(tags),
            "nodata": str(src.nodata),
            "crs": str(src.crs),
            "quality_controlled": "true",
        }
    return data, grid, metadata


def _flagged_cells_frame(result: ASOQCResult, maximum_records: int) -> pd.DataFrame:
    """Return detailed adjusted-cell records, capped to avoid huge CSV files."""
    frames: list[pd.DataFrame] = []
    remaining = max(0, int(maximum_records))
    total_adjusted = int(np.count_nonzero(result.adjusted_mask))
    for mask, reason, action in [
        (result.clipped_to_zero, "small_negative", "set_to_zero"),
        (result.masked_negative, "negative_below_tolerance", "set_to_nodata"),
        (result.masked_above_maximum, "above_hard_maximum", "set_to_nodata"),
    ]:
        if remaining <= 0:
            break
        rows, cols = np.nonzero(mask)
        take = min(remaining, len(rows))
        if take == 0:
            continue
        rows = rows[:take]
        cols = cols[:take]
        cleaned = result.cleaned_mm[rows, cols]
        frames.append(
            pd.DataFrame(
                {
                    "date": result.date,
                    "variable": result.variable,
                    "reason": reason,
                    "action": action,
                    "row": rows.astype(int),
                    "column": cols.astype(int),
                    "x": result.grid.x[cols].astype(float),
                    "y": result.grid.y[rows].astype(float),
                    "raw_value_mm": result.raw_mm[rows, cols].astype(float),
                    "cleaned_value_mm": cleaned.astype(float),
                    "inside_east_river": result.inside_watershed[rows, cols].astype(bool),
                }
            )
        )
        remaining -= take
    frame = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()
    result.summary["adjusted_cell_records_written"] = int(len(frame))
    result.summary["adjusted_cell_records_truncated"] = bool(total_adjusted > len(frame))
    return frame


def _top_values_frame(result: ASOQCResult, count: int) -> pd.DataFrame:
    finite_rows, finite_cols = np.nonzero(np.isfinite(result.raw_mm))
    if finite_rows.size == 0:
        return pd.DataFrame()
    values = result.raw_mm[finite_rows, finite_cols]
    order = np.argsort(values)[::-1][:count]
    records = []
    for rank, index in enumerate(order, start=1):
        row = int(finite_rows[index])
        col = int(finite_cols[index])
        reasons = []
        if result.clipped_to_zero[row, col]:
            reasons.append("small_negative")
        if result.masked_negative[row, col]:
            reasons.append("negative_below_tolerance")
        if result.masked_above_maximum[row, col]:
            reasons.append("above_hard_maximum")
        records.append(
            {
                "date": result.date,
                "variable": result.variable,
                "rank": rank,
                "row": row,
                "column": col,
                "x": float(result.grid.x[col]),
                "y": float(result.grid.y[row]),
                "raw_value_mm": float(result.raw_mm[row, col]),
                "inside_east_river": bool(result.inside_watershed[row, col]),
                "adjusted": bool(result.adjusted_mask[row, col]),
                "reason": ";".join(reasons),
            }
        )
    return pd.DataFrame.from_records(records)


def _overlay_boundary(ax, watershed: gpd.GeoDataFrame, grid: GridSpec) -> None:
    local = watershed.to_crs(grid.crs)
    local.boundary.plot(ax=ax, color="black", linewidth=0.8)


def _plot_qc_figure(
    cfg: WorkflowConfig,
    result: ASOQCResult,
    watershed: gpd.GeoDataFrame,
    destination: Path,
) -> Path:
    qc = cfg.section("aso")["quality_control"]
    finite_clean = result.cleaned_mm[np.isfinite(result.cleaned_mm)]
    percentile = float(qc["display_percentile"])
    vmax = float(np.percentile(finite_clean, percentile)) if finite_clean.size else 1.0
    vmax = max(vmax, 1.0)
    origin = "upper" if result.grid.transform.e < 0 else "lower"
    extent = result.grid.extent

    fig, axes = plt.subplots(2, 2, figsize=(15, 12))
    raw_image = axes[0, 0].imshow(
        result.raw_mm,
        extent=extent,
        origin=origin,
        vmin=0,
        vmax=vmax,
        cmap="viridis",
        interpolation="nearest",
    )
    axes[0, 0].set_title(f"Raw ASO (display capped at p{percentile:g})", fontweight="normal")
    _overlay_boundary(axes[0, 0], watershed, result.grid)
    fig.colorbar(raw_image, ax=axes[0, 0], fraction=0.046, pad=0.04, label="mm")

    clean_image = axes[0, 1].imshow(
        result.cleaned_mm,
        extent=extent,
        origin=origin,
        vmin=0,
        vmax=vmax,
        cmap="viridis",
        interpolation="nearest",
    )
    axes[0, 1].set_title("QC-cleaned ASO", fontweight="normal")
    _overlay_boundary(axes[0, 1], watershed, result.grid)
    fig.colorbar(clean_image, ax=axes[0, 1], fraction=0.046, pad=0.04, label="mm")

    flag_codes = np.zeros(result.grid.shape, dtype=np.uint8)
    flag_codes[result.clipped_to_zero] = 1
    flag_codes[result.masked_negative] = 2
    flag_codes[result.masked_above_maximum] = 3
    flag_plot = np.where(flag_codes > 0, flag_codes, np.nan)
    cmap = ListedColormap(["#f5d142", "#d95f02", "#d73027"])
    axes[1, 0].imshow(
        flag_plot,
        extent=extent,
        origin=origin,
        vmin=1,
        vmax=3,
        cmap=cmap,
        interpolation="nearest",
    )
    axes[1, 0].set_title("Adjusted cell locations", fontweight="normal")
    _overlay_boundary(axes[1, 0], watershed, result.grid)
    axes[1, 0].text(
        0.02,
        0.02,
        "1: set to zero   2: negative masked   3: above maximum masked",
        transform=axes[1, 0].transAxes,
        ha="left",
        va="bottom",
        fontsize=max(8, int(cfg.section("plotting")["font_size"]) - 7),
        bbox={"facecolor": "white", "alpha": 0.8, "edgecolor": "none"},
    )

    raw_hist = result.raw_mm[np.isfinite(result.raw_mm) & (result.raw_mm <= vmax)]
    clean_hist = result.cleaned_mm[np.isfinite(result.cleaned_mm) & (result.cleaned_mm <= vmax)]
    bins = np.linspace(0, vmax, 80)
    axes[1, 1].hist(raw_hist, bins=bins, histtype="step", linewidth=1.2, label="Raw")
    axes[1, 1].hist(clean_hist, bins=bins, histtype="step", linewidth=1.2, label="Cleaned")
    axes[1, 1].set_yscale("log")
    axes[1, 1].set_xlabel(f"{result.variable} (mm)")
    axes[1, 1].set_ylabel("Cell count")
    axes[1, 1].set_title("Distribution below display cap", fontweight="normal")
    axes[1, 1].legend(frameon=False)

    for ax in axes.flat[:3]:
        ax.set_xlabel("Easting (m)")
        ax.set_ylabel("Northing (m)")
        ax.ticklabel_format(style="plain", axis="both", useOffset=False)
    fig.suptitle(
        f"ASO quality control: {result.date} {result.variable}",
        fontweight="normal",
    )
    fig.text(
        0.5,
        0.01,
        (
            f"Raw max = {result.summary['raw_max_mm']:,.1f} mm; "
            f"cleaned max = {result.summary['cleaned_max_mm']:,.1f} mm; "
            f"adjusted cells = {result.summary['adjusted_count']:,}; "
            f"inside East River = {result.summary['watershed_adjusted_count']:,}."
        ),
        ha="center",
        va="bottom",
    )
    fig.tight_layout(rect=(0, 0.035, 1, 0.95))
    destination.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(destination, dpi=int(cfg.section("plotting")["dpi"]), bbox_inches="tight")
    plt.close(fig)
    return destination


def _density_summary(
    date: str,
    swe: ASOQCResult,
    depth: ASOQCResult,
    qc: dict[str, Any],
) -> dict[str, Any]:
    if swe.grid.shape != depth.grid.shape or not np.allclose(swe.grid.x, depth.grid.x) or not np.allclose(swe.grid.y, depth.grid.y):
        return {"date": date, "status": "SWE and snow-depth grids do not align"}
    minimum_depth = float(qc["density_minimum_snow_depth_mm"])
    valid = (
        np.isfinite(swe.cleaned_mm)
        & np.isfinite(depth.cleaned_mm)
        & (depth.cleaned_mm >= minimum_depth)
        & (swe.cleaned_mm >= 0)
    )
    density = np.full(swe.grid.shape, np.nan, dtype=float)
    density[valid] = 1000.0 * swe.cleaned_mm[valid] / depth.cleaned_mm[valid]
    inside = valid & swe.inside_watershed
    values = density[inside]
    values = values[np.isfinite(values)]
    if values.size == 0:
        return {"date": date, "status": "No valid density cells", "cell_count": 0}
    low = float(qc["density_review_min_kg_m3"])
    high = float(qc["density_review_max_kg_m3"])
    impossible = float(qc["density_hard_max_kg_m3"])
    p = np.percentile(values, [1, 5, 50, 95, 99])
    return {
        "date": date,
        "status": "reported_only_not_used_for_masking",
        "minimum_snow_depth_mm": minimum_depth,
        "cell_count": int(values.size),
        "density_p01_kg_m3": float(p[0]),
        "density_p05_kg_m3": float(p[1]),
        "density_p50_kg_m3": float(p[2]),
        "density_p95_kg_m3": float(p[3]),
        "density_p99_kg_m3": float(p[4]),
        "below_review_min_count": int(np.count_nonzero(values < low)),
        "above_review_max_count": int(np.count_nonzero(values > high)),
        "above_water_density_count": int(np.count_nonzero(values > impossible)),
    }


def _set_doc_font(run, size_pt: float | None = None) -> None:
    from docx.shared import Pt

    run.font.name = "Arial"
    run.bold = False
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    # Ensure East Asian font mapping also uses Arial.
    run._element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "Arial")


def _add_paragraph(document, text: str, size_pt: float = 10.5, space_after_pt: float = 5.0):
    from docx.shared import Pt

    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    _set_doc_font(run, size_pt)
    paragraph.paragraph_format.space_after = Pt(space_after_pt)
    return paragraph


def _add_heading(document, text: str, level: int = 1):
    from docx.shared import Pt

    sizes = {1: 16, 2: 13, 3: 11.5}
    paragraph = document.add_paragraph()
    paragraph.style = document.styles[f"Heading {min(level, 3)}"]
    run = paragraph.add_run(text)
    _set_doc_font(run, sizes.get(level, 11.5))
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    return paragraph


def _format_table(table, font_size: float = 8.5) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    _set_doc_font(run, font_size)


def _write_docx_report(
    cfg: WorkflowConfig,
    summary: pd.DataFrame,
    density: pd.DataFrame,
    results: list[ASOQCResult],
    figure_paths: dict[tuple[str, str], Path],
    output_path: Path,
) -> Path:
    try:
        from docx import Document
        from docx.enum.section import WD_SECTION
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt
    except ImportError as exc:  # pragma: no cover - clear runtime guidance
        raise RuntimeError(
            "python-docx is required to generate the ASO QC report. Reinstall the project "
            "environment or run: python -m pip install python-docx"
        ) from exc

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    for style_name in ["Normal", "Heading 1", "Heading 2", "Heading 3"]:
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.bold = False
    document.styles["Normal"].font.size = Pt(10.5)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("East River ASO Quality-Control Report")
    _set_doc_font(run, 20)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(datetime.now(timezone.utc).strftime("Generated %Y-%m-%d %H:%M UTC"))
    _set_doc_font(run, 10)

    _add_heading(document, "Purpose and decision rule", 1)
    _add_paragraph(
        document,
        "This report documents every automatic change made to the ASO SWE and snow-depth rasters before regridding and model comparison. The original GeoTIFFs are not modified. Cleaned copies are written in millimeters and are used by the downstream East River workflow.",
    )
    _add_paragraph(
        document,
        "The automatic policy is intentionally conservative: existing no-data values remain no-data; tiny negative values are set to zero; more-negative values and values above an explicit hard maximum are set to no-data. Percentiles are used only for display and reporting, never for clipping. Cross-variable snow-density checks are reported for review but do not automatically change either raster.",
    )

    qc = cfg.section("aso")["quality_control"]
    _add_heading(document, "Configured thresholds", 2)
    threshold_rows = [
        ("SWE hard maximum", f"{float(qc['swe_hard_max_mm']):,.0f} mm"),
        ("Snow-depth hard maximum", f"{float(qc['snow_depth_hard_max_mm']):,.0f} mm"),
        ("Small-negative tolerance", f"{float(qc['negative_clip_tolerance_mm']):g} mm"),
        ("QC map display percentile", f"p{float(qc['display_percentile']):g}"),
        ("Density diagnostic minimum snow depth", f"{float(qc['density_minimum_snow_depth_mm']):,.0f} mm"),
        ("Density review range", f"{float(qc['density_review_min_kg_m3']):g}-{float(qc['density_review_max_kg_m3']):g} kg/m3"),
        ("Water-density consistency ceiling", f"{float(qc['density_hard_max_kg_m3']):g} kg/m3"),
    ]
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Setting"
    table.rows[0].cells[1].text = "Value"
    for label, value in threshold_rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    _format_table(table, 9)

    _add_heading(document, "Raster-level summary", 1)
    table = document.add_table(rows=1, cols=8)
    table.style = "Table Grid"
    headers = ["Date", "Variable", "Raw max", "Clean max", "Adjusted", "Inside basin", "% adjusted", "Clean p99.9"]
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    for _, row in summary.iterrows():
        cells = table.add_row().cells
        values = [
            str(row["date"]),
            str(row["variable"]),
            f"{row['raw_max_mm']:,.1f}",
            f"{row['cleaned_max_mm']:,.1f}",
            f"{int(row['adjusted_count']):,}",
            f"{int(row['watershed_adjusted_count']):,}",
            f"{row['adjusted_percent_of_raw_valid']:.5f}",
            f"{row['cleaned_p99_9_mm']:,.1f}",
        ]
        for cell, value in zip(cells, values):
            cell.text = value
    _format_table(table, 7.5)

    if not density.empty:
        _add_heading(document, "SWE/snow-depth density consistency", 1)
        _add_paragraph(
            document,
            "For cells with sufficient snow depth, implied bulk snow density is calculated as 1000 x SWE / snow depth. These counts are diagnostic only because an inconsistent ratio does not identify which source raster is wrong.",
        )
        table = document.add_table(rows=1, cols=7)
        table.style = "Table Grid"
        headers = ["Date", "Cells", "p01", "p50", "p99", "Above review max", "Above 1000"]
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
        for _, row in density.iterrows():
            cells = table.add_row().cells
            values = [
                str(row.get("date", "")),
                f"{int(row.get('cell_count', 0)):,}",
                f"{row.get('density_p01_kg_m3', np.nan):,.1f}",
                f"{row.get('density_p50_kg_m3', np.nan):,.1f}",
                f"{row.get('density_p99_kg_m3', np.nan):,.1f}",
                f"{int(row.get('above_review_max_count', 0)):,}",
                f"{int(row.get('above_water_density_count', 0)):,}",
            ]
            for cell, value in zip(cells, values):
                cell.text = value
        _format_table(table, 8)

    _add_heading(document, "Per-raster diagnostics", 1)
    for index, result in enumerate(results):
        _add_heading(document, f"{result.date} - {result.variable}", 2)
        s = result.summary
        _add_paragraph(
            document,
            (
                f"Raw valid cells: {s['raw_count']:,}. Raw maximum: {s['raw_max_mm']:,.2f} mm. "
                f"Cleaned maximum: {s['cleaned_max_mm']:,.2f} mm. "
                f"Adjusted cells: {s['adjusted_count']:,} ({s['adjusted_percent_of_raw_valid']:.6f}% of raw valid cells), "
                f"including {s['watershed_adjusted_count']:,} inside the East River diagnostic mask. "
                f"Actions: {s['clipped_to_zero_count']:,} set to zero, "
                f"{s['masked_negative_count']:,} negative values masked, and "
                f"{s['masked_above_hard_max_count']:,} values above the hard maximum masked."
            ),
        )
        image_path = figure_paths[(result.date, result.variable)]
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(str(image_path), width=Inches(6.8))
        if index < len(results) - 1:
            document.add_page_break()

    _add_heading(document, "Outputs and downstream use", 1)
    _add_paragraph(
        document,
        f"QC diagnostics are stored in: {cfg.output_dir / 'diagnostics' / 'aso_qc'}",
    )
    _add_paragraph(
        document,
        f"Cleaned millimeter GeoTIFFs are stored in: {cfg.output_dir / 'cache' / 'aso_qc'}",
    )
    _add_paragraph(
        document,
        "The run-all command uses these cleaned rasters for ASO-to-iSnobal and ASO-to-Noah-MP regridding, difference maps, elevation-band distributions, and summary statistics. The original ASO files remain unchanged and can always be audited against the cleaned products.",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


def run_aso_quality_control(
    cfg: WorkflowConfig,
    watershed: gpd.GeoDataFrame,
    logger: logging.Logger,
) -> dict[str, Any]:
    """Run QC for all configured ASO dates and create machine/human reports."""
    aso = cfg.section("aso")
    qc = aso["quality_control"]
    diagnostics_dir = cfg.output_dir / "diagnostics" / "aso_qc"
    figures_dir = diagnostics_dir / "figures"
    diagnostics_dir.mkdir(parents=True, exist_ok=True)
    figures_dir.mkdir(parents=True, exist_ok=True)
    (cfg.output_dir / "cache" / "aso_qc").mkdir(parents=True, exist_ok=True)

    results: list[ASOQCResult] = []
    summaries: list[dict[str, Any]] = []
    flagged_frames: list[pd.DataFrame] = []
    top_frames: list[pd.DataFrame] = []
    figure_paths: dict[tuple[str, str], Path] = {}
    by_date: dict[str, dict[str, ASOQCResult]] = {}

    for date in aso["dates"]:
        for variable, pattern in [("SWE", aso["swe_pattern"]), ("Snow depth", aso["snow_depth_pattern"])]:
            source = Path(cfg.data["paths"]["aso_directory"]).expanduser() / pattern.format(date=date)
            raw, grid, metadata = read_aso_raster(source)
            inside = watershed_center_mask(watershed, grid)
            result = apply_aso_quality_control(
                raw,
                date=date,
                variable=variable,
                source_path=source,
                grid=grid,
                metadata=metadata,
                inside_watershed=inside,
                qc=qc,
            )
            results.append(result)
            summaries.append(result.summary)
            flagged = _flagged_cells_frame(result, int(qc["maximum_adjusted_cell_records"]))
            if not flagged.empty:
                flagged_frames.append(flagged)
            top = _top_values_frame(result, int(qc["top_values_per_raster"]))
            if not top.empty:
                top_frames.append(top)
            cleaned_path = qc_cleaned_path(cfg, date, variable)
            _write_cleaned_raster(cleaned_path, result, qc)
            figure_path = figures_dir / f"{date}_{VARIABLE_SLUGS[variable]}_qc.png"
            _plot_qc_figure(cfg, result, watershed, figure_path)
            figure_paths[(date, variable)] = figure_path
            by_date.setdefault(date, {})[variable] = result
            logger.info(
                "ASO QC %s %s: adjusted %s cells (%s inside East River); raw max %.2f mm; clean max %.2f mm",
                date,
                variable,
                result.summary["adjusted_count"],
                result.summary["watershed_adjusted_count"],
                result.summary["raw_max_mm"],
                result.summary["cleaned_max_mm"],
            )

    summary = pd.DataFrame(summaries)
    summary_path = diagnostics_dir / "aso_qc_summary.csv"
    summary.to_csv(summary_path, index=False)

    flagged_columns = [
        "date", "variable", "reason", "action", "row", "column", "x", "y",
        "raw_value_mm", "cleaned_value_mm", "inside_east_river",
    ]
    flagged_all = pd.concat(flagged_frames, ignore_index=True) if flagged_frames else pd.DataFrame(columns=flagged_columns)
    flagged_path = diagnostics_dir / "aso_qc_adjusted_cells.csv"
    flagged_all.to_csv(flagged_path, index=False)

    top_all = pd.concat(top_frames, ignore_index=True) if top_frames else pd.DataFrame()
    top_path = diagnostics_dir / "aso_qc_top_values.csv"
    top_all.to_csv(top_path, index=False)

    density_rows = []
    for date, pair in by_date.items():
        if "SWE" in pair and "Snow depth" in pair:
            density_rows.append(_density_summary(date, pair["SWE"], pair["Snow depth"], qc))
    density = pd.DataFrame(density_rows)
    density_path = diagnostics_dir / "aso_qc_density_diagnostics.csv"
    density.to_csv(density_path, index=False)

    method = {
        "generated_utc": datetime.now(timezone.utc).isoformat(),
        "principle": "Only explicit hard-invalid values are changed; percentile statistics are diagnostic only.",
        "original_files_modified": False,
        "cleaned_units": "mm",
        "quality_control_settings": qc,
        "adjusted_cell_detail_note": "Detailed adjusted-cell CSV records are capped per raster; summary counts are never truncated.",
        "actions": {
            "existing_nodata_or_nonfinite": "retain as nodata",
            "negative_within_tolerance": "set to zero",
            "negative_below_tolerance": "set to nodata",
            "above_variable_hard_maximum": "set to nodata",
            "density_outside_review_range": "report only; no automatic masking",
        },
    }
    method_path = diagnostics_dir / "aso_qc_method.json"
    save_json(method, method_path)

    report_path = diagnostics_dir / "ASO_Quality_Control_Report.docx"
    _write_docx_report(cfg, summary, density, results, figure_paths, report_path)

    manifest = {
        "report": str(report_path),
        "summary_csv": str(summary_path),
        "adjusted_cells_csv": str(flagged_path),
        "top_values_csv": str(top_path),
        "density_diagnostics_csv": str(density_path),
        "method_json": str(method_path),
        "cleaned_rasters": [str(qc_cleaned_path(cfg, result.date, result.variable)) for result in results],
        "figures": [str(figure_paths[(result.date, result.variable)]) for result in results],
    }
    manifest_path = diagnostics_dir / "aso_qc_manifest.json"
    save_json(manifest, manifest_path)
    logger.info("ASO QC report written to %s", report_path)
    return manifest
